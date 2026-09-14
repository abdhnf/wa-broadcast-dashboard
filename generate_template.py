import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_wa_contacts_template(output_path):
    wb = openpyxl.Workbook()
    
    # --- Sheet 1: Data Kontak ---
    ws = wb.active
    ws.title = "Template Kontak"
    ws.views.sheetView[0].showGridLines = True
    
    # Palet Sunset Amber & Neutral
    header_fill = PatternFill(start_color="C2610C", end_color="C2610C", fill_type="solid") # Sunset Amber
    custom_fill = PatternFill(start_color="854D0E", end_color="854D0E", fill_type="solid") # Amber Deep / Custom Var
    note_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")   # Amber Wash
    zebra_fill = PatternFill(start_color="FDFCF7", end_color="FDFCF7", fill_type="solid")
    
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=11, color="1C1917")
    mono_font = Font(name="Consolas", size=10, color="0369A1") # phone format
    note_font = Font(name="Calibri", size=10, italic=True, color="92400E")
    
    thin_border = Border(
        left=Side(style='thin', color='E2D9CC'),
        right=Side(style='thin', color='E2D9CC'),
        top=Side(style='thin', color='E2D9CC'),
        bottom=Side(style='thin', color='E2D9CC')
    )
    
    header_border = Border(
        left=Side(style='thin', color='A14C05'),
        right=Side(style='thin', color='A14C05'),
        top=Side(style='thin', color='A14C05'),
        bottom=Side(style='medium', color='78350F')
    )

    # Row 1: Banner / Petunjuk
    ws.merge_cells("A1:G1")
    banner_cell = ws["A1"]
    banner_cell.value = "📋 TEMPLATE IMPORT AUDIENS WHATSAPP BLAST — Kolom Amber = Wajib/Dasar, Kolom Cokelat = Variabel Kustom Dinamis ({{key}})"
    banner_cell.fill = note_fill
    banner_cell.font = note_font
    banner_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 24

    # Row 2: Header Kolom
    headers = [
        ("name", "Nama Kontak (Wajib)", header_fill),
        ("phone", "Nomor WhatsApp (Wajib: 628xx)", header_fill),
        ("group", "Segmen / Grup (Opsional)", header_fill),
        ("kota", "Variabel: kota", custom_fill),
        ("voucher", "Variabel: voucher", custom_fill),
        ("nominal", "Variabel: nominal", custom_fill),
        ("link", "Variabel: link", custom_fill),
    ]

    ws.row_dimensions[2].height = 28
    for col_idx, (col_key, col_title, fill) in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col_idx, value=col_key)
        cell.fill = fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = header_border

    # Sample Data (Row 3-7)
    sample_data = [
        ("Budi Santoso", "081234567890", "Pelanggan VIP", "Jakarta", "PROMO100K", "100.000", "https://srvx.id/v/budi"),
        ("Siti Rahma", "6285712345678", "Reseller Reguler", "Surabaya", "DISC25", "50.000", "https://srvx.id/v/siti"),
        ("Ahmad Fauzi", "+6281987654321", "Lead Prospek", "Bandung", "NEWUSER", "25.000", "https://srvx.id/v/ahmad"),
        ("Dewi Lestari", "085299887766", "Pelanggan VIP", "Medan", "VIP50", "150.000", "https://srvx.id/v/dewi"),
        ("Rian Pratama", "6289611223344", "Reseller Reguler", "Yogyakarta", "CASHBACK", "75.000", "https://srvx.id/v/rian"),
    ]

    for row_idx, row_values in enumerate(sample_data, 3):
        ws.row_dimensions[row_idx].height = 22
        is_zebra = (row_idx % 2 == 0)
        current_fill = zebra_fill if is_zebra else None
        
        for col_idx, val in enumerate(row_values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = mono_font if col_idx == 2 else data_font
            cell.border = thin_border
            if current_fill:
                cell.fill = current_fill
            
            # Format text explicitly agar nomor HP dengan awalan 0 tidak hilang
            cell.number_format = '@'
            if col_idx in [2, 4, 5]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx == 6:
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    # Column Widths
    col_widths = {
        'A': 22, # name
        'B': 22, # phone
        'C': 20, # group
        'D': 16, # kota
        'E': 16, # voucher
        'F': 16, # nominal
        'G': 28, # link
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    # --- Sheet 2: Petunjuk & Variabel ---
    ws2 = wb.create_sheet(title="Panduan Variabel")
    ws2.views.sheetView[0].showGridLines = True
    ws2.row_dimensions[1].height = 28
    
    title_cell = ws2["A1"]
    title_cell.value = "📖 PANDUAN PENGGUNAAN TEMPLATE & VARIABEL KUSTOM"
    title_cell.font = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
    title_cell.fill = header_fill
    title_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws2.merge_cells("A1:C1")

    guide_rows = [
        ("Kolom", "Aturan / Keterangan", "Contoh Nilai"),
        ("name", "Nama penerima pesan (Wajib diisi). Digunakan untuk variabel {{name}}.", "Budi Santoso"),
        ("phone", "Nomor WhatsApp (Wajib diisi). Format otomatis dinormalisasi ke 628xx.", "081234567890 atau 6281234567890"),
        ("group", "Nama segmen audiens / kategori kontak (Opsional). Jika kosong masuk grup default.", "Pelanggan VIP"),
        ("Variabel Kustom", "Semua kolom tambahan di samping kanan kolom group (misal: kota, voucher, tagihan, dll) otomatis disimpan sebagai Variabel Kustom Kontak.", "Kolom baru: tagihan, tenor, dsb"),
        ("Cara Panggil", "Di template pesan WA Blast, panggil menggunakan format kurung kurawal ganda.", "Halo {{name}}, voucher Anda {{voucher}} di {{kota}}"),
    ]

    for r_idx, (c1, c2, c3) in enumerate(guide_rows, 2):
        ws2.row_dimensions[r_idx].height = 24
        is_header = (r_idx == 2)
        for c_idx, val in enumerate([c1, c2, c3], 1):
            cell = ws2.cell(row=r_idx, column=c_idx, value=val)
            cell.border = thin_border
            if is_header:
                cell.font = Font(name="Calibri", size=11, bold=True, color="92400E")
                cell.fill = note_fill
            else:
                cell.font = data_font
                if c_idx == 1:
                    cell.font = Font(name="Calibri", size=11, bold=True, color="C2610C")
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    ws2.column_dimensions['A'].width = 20
    ws2.column_dimensions['B'].width = 50
    ws2.column_dimensions['C'].width = 35

    wb.save(output_path)
    print(f"Template created successfully: {output_path}")

if __name__ == "__main__":
    generate_wa_contacts_template("/home/abdhnf/projects/wa-broadcast-dashboard/public/template-kontak-blast.xlsx")
